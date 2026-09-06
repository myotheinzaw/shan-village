import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SAND   = RGBColor(0xFA, 0xF5, 0xEA)
INK    = RGBColor(0x24, 0x1C, 0x15)
MUTED  = RGBColor(0x75, 0x6B, 0x58)
FAINT  = RGBColor(0x9C, 0x91, 0x7A)
ORANGE = RGBColor(0xC1, 0x50, 0x1F)
DEEP   = RGBColor(0x9C, 0x2B, 0x1F)
DARK   = RGBColor(0x17, 0x13, 0x10)
GOLD   = RGBColor(0xE8, 0xC7, 0x7A)
CREAM  = RGBColor(0xF7, 0xEF, 0xDD)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xE3, 0xD7, 0xBE)
SUNK   = RGBColor(0xF3, 0xEC, 0xDC)
GREEN  = RGBColor(0x2F, 0x7D, 0x4F)

DISP, BODY, MONO = "Georgia", "Calibri", "Consolas"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, lw=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def rect(s, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for item in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        if isinstance(item, tuple):
            item = [item]
        for t, sz, col, bold, font in item:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col
            r.font.bold = bold; r.font.name = font
    return tb


def bg(s):
    rect(s, 0, 0, W, H, SAND)


def header(s, kicker, title, sub=None):
    bg(s)
    rect(s, 0, 0, W, Inches(1.35), DARK)
    rect(s, 0, Inches(1.35), W, Pt(3), ORANGE)
    text(s, Inches(0.7), Inches(0.3), Inches(11.9), Inches(0.28),
         [(kicker.upper(), 10.5, GOLD, True, BODY)])
    text(s, Inches(0.7), Inches(0.6), Inches(11.9), Inches(0.55),
         [(title, 27, CREAM, True, DISP)])
    if sub:
        text(s, Inches(0.7), Inches(1.62), Inches(11.9), Inches(0.3),
             [(sub, 13, MUTED, False, BODY)])


def foot(s, n):
    text(s, Inches(0.7), Inches(6.95), Inches(8), Inches(0.25),
         [("Shan Village · Operations Management System", 9, FAINT, False, BODY)])
    text(s, Inches(11.2), Inches(6.95), Inches(1.43), Inches(0.25),
         [(str(n), 9, FAINT, False, BODY)], align=PP_ALIGN.RIGHT)


